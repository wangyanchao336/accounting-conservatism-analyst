#!/usr/bin/env python3
"""
汇总多家公司的财务稳健性分析结果，生成综合对比表。

用法：
  python aggregate_results.py <JSON文件目录> [输出格式: xlsx|csv|md]

示例：
  python aggregate_results.py ./results/ xlsx   # 生成 Excel 汇总表
  python aggregate_results.py ./results/ csv    # 生成 CSV
  python aggregate_results.py ./results/ md     # 生成 Markdown 表格
"""

import json
import os
import sys
import glob
from datetime import datetime

DIMENSIONS = [
    "研发费用",
    "商业信用与收入",
    "坏账计提",
    "固定资产折旧",
    "存货",
    "资产减值",
    "管理层基调",
    "其他警示信号",
]


def extract_judgment(detail_text):
    """从详细分析文本中提取判断标签（稳健/中性/不稳健）"""
    if not detail_text:
        return "-"
    if "不稳健" in detail_text:
        return "不稳健"
    if "稳健" in detail_text:
        return "稳健"
    if "中性" in detail_text:
        return "中性"
    return "-"


def extract_short_reason(detail_text, max_len=50):
    """提取简短理由"""
    if not detail_text:
        return ""
    # 取第一句或前 max_len 字
    for sep in ["。", "；", "，"]:
        if sep in detail_text:
            detail_text = detail_text.split(sep)[0]
            break
    if len(detail_text) > max_len:
        detail_text = detail_text[:max_len] + "…"
    return detail_text


def load_analysis(filepath):
    """加载单个分析 JSON 文件"""
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    row = {
        "股票代码": data.get("股票代码", "-"),
        "公司简称": data.get("公司简称", "-"),
        "报告年份": data.get("报告年份", "-"),
        "分类结果": data.get("分类结果", "-"),
        "判断置信度": data.get("判断置信度", "-"),
    }

    detail = data.get("详细分析", {})
    for dim in DIMENSIONS:
        row[f"{dim}_判断"] = extract_judgment(detail.get(dim, ""))
        row[f"{dim}_摘要"] = extract_short_reason(detail.get(dim, ""))

    row["综合评语"] = (data.get("综合评语") or "")[:120]
    return row


def count_judgments(row, judgment):
    """统计某一判断的出现次数"""
    return sum(
        1 for dim in DIMENSIONS if row.get(f"{dim}_判断") == judgment
    )


def generate_markdown(rows):
    """生成 Markdown 表格"""
    # 汇总统计
    n_robust = sum(1 for r in rows if r["分类结果"] == "稳健")
    n_neutral = sum(1 for r in rows if r["分类结果"] == "中性")
    n_aggressive = sum(1 for r in rows if r["分类结果"] == "不稳健")

    header = """# 财务报告稳健性分析 — 综合对比表

生成时间：{time}
公司数量：{total}（稳健 {r} / 中性 {n} / 不稳健 {a}）

""".format(
        time=datetime.now().strftime("%Y-%m-%d %H:%M"),
        total=len(rows),
        r=n_robust,
        n=n_neutral,
        a=n_aggressive,
    )

    # 简表：核心字段
    table = "| 代码 | 简称 | 年份 | 分类 | 置信度 | 研发 | 信用 | 坏账 | 折旧 | 存货 | 减值 | 基调 | 警示 |\n"
    table += "|------|------|------|------|--------|------|------|------|------|------|------|------|------|\n"

    abbrev = {
        "稳健": "✅",
        "中性": "⚠️",
        "不稳健": "❌",
        "-": "—",
    }

    for row in rows:
        table += (
            f"| {row['股票代码']} "
            f"| {row['公司简称']} "
            f"| {row['报告年份']} "
            f"| **{row['分类结果']}** "
            f"| {row['判断置信度']} "
        )
        for dim in DIMENSIONS:
            j = row[f"{dim}_判断"]
            table += f"| {abbrev.get(j, j)} "
        table += "|\n"

    # 综合评语
    remarks = "\n## 综合评语\n\n"
    for row in rows:
        if row.get("综合评语"):
            remarks += f"- **{row['公司简称']}**：{row['综合评语']}\n"

    return header + table + remarks


def generate_csv(rows):
    """生成 CSV"""
    import csv
    import io

    output = io.StringIO()
    fieldnames = [
        "股票代码", "公司简称", "报告年份", "分类结果", "判断置信度"
    ] + [f"{dim}_判断" for dim in DIMENSIONS] + ["综合评语"]

    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

    return output.getvalue()


def generate_xlsx(rows, output_path):
    """生成 Excel 文件（含样式）"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("需要 openpyxl 库。请运行: pip install openpyxl")
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "稳健性分析汇总"

    # 颜色定义
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    robust_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    neutral_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    aggressive_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    normal_font = Font(name="微软雅黑", size=10)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 表头
    headers = (
        ["股票代码", "公司简称", "报告年份", "分类结果", "置信度"]
        + ["研发费用", "商业信用", "坏账计提", "固定资产折旧", "存货", "资产减值", "管理层基调"]
        + ["警示信号摘要", "综合评语"]
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 数据行
    for row_idx, row in enumerate(rows, 2):
        values = [
            row["股票代码"], row["公司简称"], row["报告年份"],
            row["分类结果"], row["判断置信度"],
        ]
        for dim in DIMENSIONS[:7]:
            values.append(row[f"{dim}_判断"])
        # 警示信号摘要
        signals = row.get("其他警示信号_摘要", "")
        if signals and "未发现" not in signals:
            signals_text = signals[:60]
        else:
            signals_text = "无"
        values.append(signals_text)
        values.append(row.get("综合评语", "")[:200])

        classification = row["分类结果"]
        if "稳健" in classification:
            row_fill = robust_fill
        elif "不稳健" in classification:
            row_fill = aggressive_fill
        else:
            row_fill = neutral_fill

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = normal_font
            fill = robust_fill if "稳健" in classification else (
                aggressive_fill if "不稳健" in classification else neutral_fill
            )
            # 维度列按各自判断着色
            if col_idx >= 6 and col_idx <= 12:
                dim_judgment = value
                if dim_judgment == "稳健":
                    cell.fill = robust_fill
                elif dim_judgment == "不稳健":
                    cell.fill = aggressive_fill
                elif dim_judgment == "中性":
                    cell.fill = neutral_fill
            else:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    # 列宽
    widths = [12, 14, 8, 10, 8] + [11] * 7 + [20, 50]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 添加统计 sheet
    ws2 = wb.create_sheet("统计")
    ws2.cell(row=1, column=1, value="分类").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="数量").font = Font(bold=True)
    ws2.cell(row=2, column=1, value="稳健")
    ws2.cell(row=2, column=2, value=sum(1 for r in rows if r["分类结果"] == "稳健"))
    ws2.cell(row=3, column=1, value="中性")
    ws2.cell(row=3, column=2, value=sum(1 for r in rows if r["分类结果"] == "中性"))
    ws2.cell(row=4, column=1, value="不稳健")
    ws2.cell(row=4, column=2, value=sum(1 for r in rows if r["分类结果"] == "不稳健"))

    # 各维度统计
    for dim_idx, dim in enumerate(DIMENSIONS[:7]):
        col = dim_idx + 4
        ws2.cell(row=1, column=col, value=dim).font = Font(bold=True)
        n_robust = sum(1 for r in rows if r[f"{dim}_判断"] == "稳健")
        n_neutral = sum(1 for r in rows if r[f"{dim}_判断"] == "中性")
        n_agg = sum(1 for r in rows if r[f"{dim}_判断"] == "不稳健")
        total = len(rows)
        ws2.cell(row=2, column=col, value=f"稳健 {n_robust} ({n_robust/total*100:.0f}%)")
        ws2.cell(row=3, column=col, value=f"中性 {n_neutral} ({n_neutral/total*100:.0f}%)")
        ws2.cell(row=4, column=col, value=f"不稳健 {n_agg} ({n_agg/total*100:.0f}%)")

    wb.save(output_path)
    return True


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    json_dir = sys.argv[1]
    fmt = sys.argv[2] if len(sys.argv) > 2 else "md"

    # 收集所有 JSON 文件
    json_files = glob.glob(os.path.join(json_dir, "*.json"))
    if not json_files:
        print(f"错误：在 {json_dir} 中未找到 JSON 文件")
        sys.exit(1)

    rows = []
    for f in sorted(json_files):
        try:
            row = load_analysis(f)
            rows.append(row)
        except Exception as e:
            print(f"跳过 {f}: {e}")

    if not rows:
        print("错误：没有成功加载任何分析结果")
        sys.exit(1)

    # 按分类结果排序（不稳健在前）
    sort_order = {"不稳健": 0, "中性": 1, "稳健": 2}
    rows.sort(key=lambda r: sort_order.get(r["分类结果"], 9))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if fmt == "xlsx":
        output_path = os.path.join(json_dir, f"汇总表_{timestamp}.xlsx")
        if generate_xlsx(rows, output_path):
            print(f"Excel 汇总表已生成: {output_path}")
    elif fmt == "csv":
        csv_content = generate_csv(rows)
        output_path = os.path.join(json_dir, f"汇总表_{timestamp}.csv")
        with open(output_path, "w", encoding="utf-8-sig") as f:
            f.write(csv_content)
        print(f"CSV 汇总表已生成: {output_path}")
    elif fmt == "md":
        md_content = generate_markdown(rows)
        output_path = os.path.join(json_dir, f"汇总表_{timestamp}.md")
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        print(f"Markdown 汇总表已生成: {output_path}")
    else:
        print(f"不支持的格式: {fmt}，可选: xlsx, csv, md")
        sys.exit(1)


if __name__ == "__main__":
    main()
